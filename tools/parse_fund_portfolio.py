#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""تبدیل پیوندهای Excel پرتفوی ماهانهٔ صندوق‌ها (کدال) به data/funds/holdings_<نماد>.csv

    python tools/parse_fund_portfolio.py

ورودی: data/raw/codal/funds/<نماد>/L_<TracingNo>.meta.json و L_<TracingNo>_att*.xlsx (برگهٔ «سهام»: صورت وضعیت پرتفوی).
خروجی: date (پایان ماه، ISO)، symbol (نماد TSETMC یا خالی)، name_fa (نام در کدال)، weight_pct (درصد از کل دارایی‌ها)،
       qty، price، matched. نمادهای تطبیق‌نیافته در data/funds/holdings_unmatched.csv.
"""
import csv
import difflib
import re
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from fetch_index_funds import j2g, norm  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "codal" / "funds"
OUT = ROOT / "data" / "funds"
FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def key(s: str) -> str:
    s = norm(s).replace("‌", " ")
    s = re.sub(r"[\s\.\-\(\)،,_]+", "", s)
    return s


def load_names() -> dict:
    """نام کامل و نماد هر سهم (instruments.csv، instruments_extra.csv) → نماد. فقط سهام."""
    out = {}
    for fn in ("instruments.csv", "instruments_extra.csv"):
        p = ROOT / "data" / fn
        if not p.exists():
            continue
        for r in csv.DictReader(p.open(encoding="utf-8")):
            if r.get("type", "stock") not in ("stock", "fund", ""):
                continue
            sym = norm(r["symbol"])
            if re.search(r"\d$", sym):
                continue
            out.setdefault(key(r["name"]), sym)
            out.setdefault(key(sym), sym)
    return out


# نام کدال که با نام TSETMC فرق دارد (با خروجی holdings_unmatched.csv تکمیل شود)
ALIASES = {"سنگ آهن گهرزمین": "کگهر"}


def match(name: str, names: dict) -> str:
    k = key(name)
    for a, sym in ALIASES.items():
        if key(a) == k:
            return sym
    if k in names:
        return names[k]
    if name.strip().startswith(("ح .", "ح.", "ح ")):
        return ""                                              # حق تقدم
    cand = difflib.get_close_matches(k, names.keys(), n=1, cutoff=0.9)
    return names[cand[0]] if cand else ""


def month_end(title: str):
    m = re.search(r"منتهی\s*به\s*([\d۰-۹]{4})/([\d۰-۹]{1,2})/([\d۰-۹]{1,2})", title.translate(FA).replace("‌", " "))
    if not m:
        return None
    y, mo, d = (int(x) for x in m.groups())
    return j2g(y, mo, d).isoformat()


def parse_xlsx(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["سهام"] if "سهام" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r and r[0] and "نام شرکت" in str(r[0])), None)
    if hdr is None:
        return []
    h = rows[hdr]
    last = lambda lab: max((j for j, v in enumerate(h) if v and lab in str(v)), default=None)
    c_pct, c_price, c_qty = last("درصد به کل"), last("قیمت بازار"), last("تعداد")
    out = []
    for r in rows[hdr + 1:]:
        if not r or r[0] is None:
            continue
        nm = str(r[0]).strip()
        if nm.startswith("جمع"):
            break
        v = r[c_pct] if c_pct is not None else None
        if isinstance(v, (int, float)):
            out.append({"name_fa": nm, "weight_pct": float(v), "qty": r[c_qty] if c_qty is not None else None,
                        "price": r[c_price] if c_price is not None else None})
    return out


def main():
    names = load_names()
    OUT.mkdir(parents=True, exist_ok=True)
    unmatched = {}
    for fdir in sorted(p for p in RAW.iterdir() if p.is_dir()):
        sym = fdir.name
        recs = []
        for mp in sorted(fdir.glob("L_*.meta.json")):
            meta = json.loads(mp.read_text(encoding="utf-8"))
            title = meta.get("Title") or ""
            if "پورتفوی" not in title and "پرتفوی" not in title:
                continue
            d = month_end(title)
            xs = sorted(fdir.glob(mp.name.replace(".meta.json", "_att*.xls*")))
            if not d or not xs:
                continue
            for it in parse_xlsx(xs[0]):
                s = match(it["name_fa"], names)
                if not s:
                    unmatched[(sym, it["name_fa"])] = unmatched.get((sym, it["name_fa"]), 0) + 1
                recs.append({"date": d, "symbol": s, "name_fa": it["name_fa"], "weight_pct": round(it["weight_pct"], 4),
                             "qty": it["qty"], "price": it["price"], "matched": int(bool(s))})
        if not recs:
            continue
        recs.sort(key=lambda r: (r["date"], -r["weight_pct"]))
        with (OUT / f"holdings_{sym}.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(recs[0].keys()))
            w.writeheader()
            w.writerows(recs)
        ds = sorted({r["date"] for r in recs})
        tot = {d: sum(r["weight_pct"] for r in recs if r["date"] == d) for d in ds}
        mt = {d: sum(r["weight_pct"] for r in recs if r["date"] == d and r["matched"]) for d in ds}
        print(f"{sym}: {len(ds)} ماه ({ds[0]} تا {ds[-1]})؛ میانگین سهام {sum(tot.values())/len(ds):.1f}٪ از دارایی؛ تطبیق‌یافته {sum(mt.values())/sum(tot.values())*100:.1f}٪؛ "
              f"میانگین تعداد نماد {len(recs)/len(ds):.0f}")
    if unmatched:
        with (OUT / "holdings_unmatched.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["fund", "name_fa", "months"])
            for (f, n), c in sorted(unmatched.items(), key=lambda kv: (-kv[1], kv[0])):
                w.writerow([f, n, c])
        print(f"تطبیق‌نیافته: {len(unmatched)} (فایل holdings_unmatched.csv)")


if __name__ == "__main__":
    main()
